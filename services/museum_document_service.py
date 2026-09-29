import io
from typing import Any, Dict, List
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

from utils.logger import logger


def _apply_run_font(run, *, is_bold: bool = False, font_size: int = 14) -> None:
    """Гарантує застосування шрифту Times New Roman (включно з кирилицею в OpenXML) для run."""
    run.bold = is_bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(font_size)
    run.font.color.rgb = RGBColor(0, 0, 0)

    r_pr = run._r.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), "Times New Roman")
    r_fonts.set(qn("w:hAnsi"), "Times New Roman")
    r_fonts.set(qn("w:cs"), "Times New Roman")


def _set_cell_style(
    cell,
    text: str,
    *,
    is_bold: bool = False,
    font_size: int = 14,
    alignment: WD_ALIGN_PARAGRAPH = WD_ALIGN_PARAGRAPH.LEFT
) -> None:
    """Встановлює уніфікований стиль для комірки таблиці: Times New Roman 14pt, інтервал 1.15, вертикальне центрування."""
    cell.text = text
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    paragraph = cell.paragraphs[0]
    paragraph.alignment = alignment
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)

    for run in paragraph.runs:
        _apply_run_font(run, is_bold=is_bold, font_size=font_size)


def _apply_table_borders(table) -> None:
    """Накладає суцільну тонку чорну рамку на всі межі комірок таблиці."""
    tbl_pr = table._tbl.tblPr
    borders_element = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'  <w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'  <w:insideV w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f"</w:tblBorders>"
    )
    tbl_pr.append(borders_element)


def _configure_table_layout(table, column_widths: List[Inches]) -> None:
    """
    Фіксує макет таблиці (w:tblLayout type="fixed"), компактні внутрішні відступи комірок (w:tblCellMar)
    та синхронізує ширини колонок у w:tblGrid для точного відображення у Word, LibreOffice та Google Docs.
    """
    tbl_pr = table._tbl.tblPr

    tbl_layout = parse_xml(f'<w:tblLayout {nsdecls("w")} w:type="fixed"/>')
    tbl_pr.append(tbl_layout)

    cell_margins = parse_xml(
        f'<w:tblCellMar {nsdecls("w")}>'
        f'  <w:top w:w="72" w:type="dxa"/>'
        f'  <w:left w:w="80" w:type="dxa"/>'
        f'  <w:bottom w:w="72" w:type="dxa"/>'
        f'  <w:right w:w="80" w:type="dxa"/>'
        f"</w:tblCellMar>"
    )
    tbl_pr.append(cell_margins)

    for idx, width in enumerate(column_widths):
        table.columns[idx].width = width

    tbl_grid = table._tbl.tblGrid
    if tbl_grid is not None:
        grid_cols = tbl_grid.findall(qn("w:gridCol"))
        for idx, width in enumerate(column_widths):
            dxa_val = str(int(round(width.inches * 1440)))
            if idx < len(grid_cols):
                grid_cols[idx].set(qn("w:w"), dxa_val)


def _apply_row_rules(row, *, is_header: bool = False) -> None:
    """Запобігає розриву рядка при переносі на нову сторінку, а для шапки включає повтор на кожній сторінці."""
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
    if is_header:
        tr_pr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))


def generate_visitors_document(payload: Dict[str, Any]) -> Dict[str, Any]:
    """
    Генерує офіційний документ DOCX зі списком відвідувачів на обрану дату екскурсії.
    Слідує патерну RORO: приймає словник payload, повертає словник з результатом.

    Параметри payload:
      - excursion_date (str): дата екскурсії (наприклад, "26.09.2026")
      - excursion_type (str): "holiday" або "regular"
      - visitors (List[Dict[str, str]]): список розгорнутих записів відвідувачів:
            [{"reg_date": "...", "name": "...", "phone": "..."}]

    Повертає:
      - stream (io.BytesIO): бінарний потік згенерованого .docx файлу
      - filename (str): рекомендована назва файлу
      - total_count (int): загальна кількість відвідувачів у документі
    """
    excursion_date: str = payload.get("excursion_date", "").strip()
    excursion_type: str = payload.get("excursion_type", "regular").strip()
    visitors: List[Dict[str, str]] = payload.get("visitors", [])

    if not excursion_date:
        raise ValueError("excursion_date must not be empty")

    is_holiday: bool = excursion_type == "holiday"
    type_label: str = "СВЯТКОВУ" if is_holiday else "СТАНДАРТНУ"

    logger.info(
        f"📄 Початок генерації документу відвідувачів ({type_label}) на дату {excursion_date}, "
        f"кількість записів: {len(visitors)}"
    )

    doc = Document()

    # 1. Альбомна орієнтація (Landscape A4: 297 x 210 мм) та оптимізовані поля 0.4" (10.16 мм)
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11.69)
    section.page_height = Inches(8.27)
    section.left_margin = Inches(0.4)
    section.right_margin = Inches(0.4)
    section.top_margin = Inches(0.4)
    section.bottom_margin = Inches(0.4)

    # 2. Налаштування базового стилю Normal
    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(14)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # 3. Шапка документа зверху таблиці
    header_lines = [
        "СПИСОК ЗАРЕЄСТРОВАНИХ ВІДВІДУВАЧІВ",
        f"НА {type_label} ЕКСКУРСІЮ",
        "ДО МУЗЕЮ КП “ОМЕТ”",
        "(м. Одеса, пл. Олексіївська, 1А)",
        f"{excursion_date}",
    ]

    for line in header_lines:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.line_spacing = 1.15
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(2)

        run = paragraph.add_run(line)
        _apply_run_font(run, is_bold=True, font_size=14)

    # Компактний відступ перед таблицею
    empty_p = doc.add_paragraph()
    empty_p.paragraph_format.space_before = Pt(0)
    empty_p.paragraph_format.space_after = Pt(4)

    # 4. Колонки таблиці
    column_headers = [
        "№ п/п",
        "Дата реєстрації на екскурсію",
        "П.І.Б. зареєстрованих відвідувачів",
        "Контактний номер телефону",
        "Згода на гарантування дотримання правил з техніки безпеки та поведінки на території підвищеної небезпеки",
    ]

    # Оптимізовані ширини колонок для альбомного формату A4 (доступна ширина 10.89 дюймів = 276.6 мм)
    column_widths = [
        Inches(0.55),  # № п/п (~14.0 мм — вміщує "№ п/п" та двозначні номери в 1 рядок)
        Inches(1.65),  # Дата реєстрації (~41.9 мм — вміщує "ДД.ММ.РРРР ГГ:ХХ" в 1 рядок)
        Inches(3.65),  # ПІБ (~92.7 мм — максимальний простір для довгих ПІБ та віку)
        Inches(1.54),  # Телефон (~39.1 мм — вміщує номер телефону в 1 рядок без переносу)
        Inches(3.50),  # Згода / підпис ТБ (~88.9 мм — компактна шапка та широке поле для підпису)
    ]

    data_alignments = [
        WD_ALIGN_PARAGRAPH.CENTER,  # № п/п
        WD_ALIGN_PARAGRAPH.CENTER,  # Дата реєстрації
        WD_ALIGN_PARAGRAPH.LEFT,    # ПІБ (без розтягування пробілів між прізвищем та ім'ям)
        WD_ALIGN_PARAGRAPH.CENTER,  # Контактний номер телефону
        WD_ALIGN_PARAGRAPH.CENTER,  # Згода / підпис
    ]

    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    _apply_table_borders(table)
    _configure_table_layout(table, column_widths)

    # Шапка таблиці
    header_row = table.rows[0]
    _apply_row_rules(header_row, is_header=True)

    for idx, header_text in enumerate(column_headers):
        cell = header_row.cells[idx]
        cell.width = column_widths[idx]
        _set_cell_style(
            cell,
            header_text,
            is_bold=True,
            font_size=14,
            alignment=WD_ALIGN_PARAGRAPH.CENTER,
        )

    # 5. Заповнення рядків відвідувачів
    for visitor_idx, visitor in enumerate(visitors, start=1):
        row = table.add_row()
        _apply_row_rules(row, is_header=False)

        reg_date_text = str(visitor.get("reg_date", "")).strip()
        name_text = str(visitor.get("name", "")).strip()
        phone_text = str(visitor.get("phone", "")).strip()
        signature_placeholder = ""  # Порожнє місце під власноручний підпис

        cell_values = [
            str(visitor_idx),
            reg_date_text,
            name_text,
            phone_text,
            signature_placeholder,
        ]

        for col_idx, value in enumerate(cell_values):
            cell = row.cells[col_idx]
            cell.width = column_widths[col_idx]
            _set_cell_style(
                cell,
                value,
                is_bold=False,
                font_size=14,
                alignment=data_alignments[col_idx],
            )

    # 6. Підсумковий блок
    doc.add_paragraph().paragraph_format.space_before = Pt(6)

    footer_summary = doc.add_paragraph()
    footer_summary.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    footer_summary.paragraph_format.line_spacing = 1.15
    footer_summary.paragraph_format.space_before = Pt(0)
    footer_summary.paragraph_format.space_after = Pt(4)

    run_summary = footer_summary.add_run(f"Всього зареєстровано відвідувачів: {len(visitors)} осіб.")
    _apply_run_font(run_summary, is_bold=True, font_size=14)

    footer_sign = doc.add_paragraph()
    footer_sign.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    footer_sign.paragraph_format.line_spacing = 1.15
    footer_sign.paragraph_format.space_before = Pt(4)
    footer_sign.paragraph_format.space_after = Pt(0)

    run_sign = footer_sign.add_run(
        "Відповідальний за проведення інструктажу з ТБ: ___________________ / ___________________ /"
    )
    _apply_run_font(run_sign, is_bold=False, font_size=14)

    # 7. Збереження у бінарний потік BytesIO
    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)

    sanitized_date = excursion_date.replace(".", "_").replace(":", "_").replace(" ", "_")
    type_suffix = "святкова" if is_holiday else "стандартна"
    filename = f"Список_відвідувачів_Музей_{type_suffix}_{sanitized_date}.docx"

    logger.info(f"✅ Документ успішно сформовано: {filename} ({len(visitors)} відвідувачів)")

    return {
        "stream": stream,
        "filename": filename,
        "total_count": len(visitors),
    }

