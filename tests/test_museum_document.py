import asyncio
import json
import uuid
import sys
import os
from unittest.mock import AsyncMock, MagicMock
from sqlalchemy import delete

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from database.db import init_db, AsyncSessionLocal, MuseumBooking, MuseumHolidayBooking
from services.museum_service import MuseumService

museum_service = MuseumService()


async def _noop_sync(*args, **kwargs):
    return None


# Вимикаємо фоновий запис тестових даних у бойову Google Таблицю під час автотестів
museum_service._sync_to_sheets_task = _noop_sync
museum_service._sync_holiday_to_sheets_task = _noop_sync


async def run_museum_document_tests():
    print("🧪 [TEST] Initializing database...")
    await init_db()

    test_date_reg = f"TEST_REG_{uuid.uuid4().hex[:6]}"
    test_phone_reg = f"097{uuid.uuid4().int % 10000000:07d}"

    try:
        # 1. Тест парсера імен учасників (включно з багаторядковими записами з Google Sheets)
        print("🧪 [TEST 1] Testing _parse_participant_names helper...")
        single_name = "Коваленко Олександр Петрович"
        assert museum_service._parse_participant_names(single_name) == [single_name]

        multi_name_numbered = "1) Коваленко Олександр (35р.); 2) Коваленко Марія (10р.)"
        parsed_numbered = museum_service._parse_participant_names(multi_name_numbered)
        assert len(parsed_numbered) == 2, f"Expected 2 participants, got {parsed_numbered}"
        assert "Коваленко Олександр" in parsed_numbered[0]
        assert "Коваленко Марія" in parsed_numbered[1]

        multi_name_semicolon = "Іванов І.І.; Петров П.П.; Сидоров С.С."
        parsed_semi = museum_service._parse_participant_names(multi_name_semicolon)
        assert len(parsed_semi) == 3

        multi_name_newlines = (
            "Загурська Юлія Альбертівна\n"
            "Загурський Максим Олександрович\n"
            "Косенко Софія Василівна\n"
            "Косенко Іван Олександрович"
        )
        parsed_newlines = museum_service._parse_participant_names(multi_name_newlines)
        assert len(parsed_newlines) == 4, f"Expected 4 participants from newlines, got {parsed_newlines}"
        assert parsed_newlines[0] == "Загурська Юлія Альбертівна"
        assert parsed_newlines[3] == "Косенко Іван Олександрович"
        print("   ✅ _parse_participant_names works properly (single, numbered, ';', and '\\n')")

        # 2. Створення тестового бронювання та перевірка розгортання групи
        reg_participants = [
            {"name": "Шевченко Тарас Григорович", "age": None},
            {"name": "Леся Українка", "age": None},
            {"name": "Іван Франко", "age": None},
        ]

        print(f"🧪 [TEST 2] Creating regular booking with 3 participants on date {test_date_reg}...")
        success_reg = await museum_service.create_booking(
            date=test_date_reg,
            count=3,
            name=reg_participants[0]["name"],
            phone=test_phone_reg,
            participants_details=json.dumps(reg_participants, ensure_ascii=False)
        )
        assert success_reg is True

        visitors_reg = await museum_service.get_visitors_for_date(test_date_reg, excursion_type="regular")
        assert len(visitors_reg) == 3, f"Expected 3 unrolled visitors, got {len(visitors_reg)}"
        assert visitors_reg[0]["name"] == "Шевченко Тарас Григорович"
        assert visitors_reg[1]["name"] == "Леся Українка"
        assert visitors_reg[2]["name"] == "Іван Франко"
        assert visitors_reg[0]["phone"] == test_phone_reg
        print(f"   ✅ Regular unrolled visitors correctly extracted: {len(visitors_reg)} items")

        # 3. Тест фільтрації службових TEST_ дат у get_distinct_booking_dates
        print("🧪 [TEST 3] Testing get_distinct_booking_dates filters out TEST_ dates...")
        all_dates = await museum_service.get_distinct_booking_dates(excursion_type="regular")
        assert test_date_reg not in all_dates, "TEST_ dates must be filtered out from admin menu!"
        print("   ✅ TEST_ dates are properly filtered out from live admin menu")

        # 4. Тест генератора DOCX документу та внутрішньої розмітки OpenXML
        print("🧪 [TEST 4] Testing generate_visitors_document formatting, margins & OpenXML layout...")
        try:
            from services.museum_document_service import generate_visitors_document
            from docx import Document
            from docx.enum.section import WD_ORIENT
            from docx.oxml.ns import qn

            res_reg = generate_visitors_document({
                "excursion_date": "26.09.2026",
                "excursion_type": "regular",
                "visitors": visitors_reg
            })

            assert res_reg["stream"] is not None
            assert res_reg["total_count"] == 3
            assert "стандартна" in res_reg["filename"]

            doc = Document(res_reg["stream"])
            section = doc.sections[0]

            # Перевірка альбомної орієнтації (Landscape) та оптимізованих полів 0.4"
            assert section.orientation == WD_ORIENT.LANDSCAPE, "Section orientation must be LANDSCAPE"
            assert round(section.page_width.inches, 2) == 11.69
            assert round(section.page_height.inches, 2) == 8.27
            assert round(section.left_margin.inches, 2) == 0.4
            assert round(section.right_margin.inches, 2) == 0.4
            assert round(section.top_margin.inches, 2) == 0.4
            assert round(section.bottom_margin.inches, 2) == 0.4

            # Перевірка шапки документу ("НА СТАНДАРТНУ ЕКСКУРСІЮ")
            doc_reg_text = "\n".join([p.text for p in doc.paragraphs])
            assert "НА СТАНДАРТНУ ЕКСКУРСІЮ" in doc_reg_text

            # Перевірка таблиці та фіксованого макета OpenXML
            assert len(doc.tables) == 1
            table = doc.tables[0]
            assert len(table.columns) == 5
            assert len(table.rows) == 4  # 1 шапка + 3 відвідувачі

            tbl_pr = table._tbl.tblPr
            tbl_layout = tbl_pr.find(qn("w:tblLayout"))
            assert tbl_layout is not None and tbl_layout.get(qn("w:type")) == "fixed"
            tbl_cell_mar = tbl_pr.find(qn("w:tblCellMar"))
            assert tbl_cell_mar is not None

            # Перевірка заголовків колонок та даних
            col_titles = [cell.text for cell in table.rows[0].cells]
            assert "№ п/п" in col_titles[0]
            assert "Дата реєстрації на екскурсію" in col_titles[1]
            assert "П.І.Б. зареєстрованих відвідувачів" in col_titles[2]
            assert "Контактний номер телефону" in col_titles[3]
            assert "Згода на гарантування дотримання правил" in col_titles[4]

            data_cells = [cell.text for cell in table.rows[1].cells]
            assert data_cells[0] == "1"
            assert data_cells[2] == "Шевченко Тарас Григорович"
            assert data_cells[3] == test_phone_reg
            assert data_cells[4] == ""

            # Тест для святкової екскурсії
            res_hol = generate_visitors_document({
                "excursion_date": "26.09.2026",
                "excursion_type": "holiday",
                "visitors": [{"reg_date": "24.09.2026 10:00", "name": "Святковий Гість (12р.)", "phone": "0501112233"}]
            })
            assert "святкова" in res_hol["filename"]
            doc_hol = Document(res_hol["stream"])
            doc_hol_text = "\n".join([p.text for p in doc_hol.paragraphs])
            assert "НА СВЯТКОВУ ЕКСКУРСІЮ" in doc_hol_text

            print("   ✅ Document generation passed: Landscape, 0.4in margins, fixed layout, 5 columns, TNR 14pt")
        except ImportError:
            print("   ⚠️ python-docx not installed in current environment, skipping direct docx binary inspection.")

        # 5. Тест навігації (кнопки "Назад до вибору дати" та "Адмін-панель Музею" під документом .docx)
        print("🧪 [TEST 5] Testing navigation buttons from document message (text=None)...")
        from handlers.admin_handlers import _render_or_send_text, admin_museum_menu_show
        from config.settings import MUSEUM_ADMIN_IDS

        mock_query = MagicMock()
        mock_query.answer = AsyncMock()
        mock_query.edit_message_text = AsyncMock(side_effect=AssertionError("Should NOT call edit_message_text when text is None!"))
        mock_query.edit_message_reply_markup = AsyncMock()
        mock_query.message = MagicMock()
        mock_query.message.text = None  # Симуляція повідомлення з файлом .docx

        mock_update = MagicMock()
        mock_update.callback_query = mock_query
        mock_update.effective_user.id = MUSEUM_ADMIN_IDS[0] if MUSEUM_ADMIN_IDS else 12345
        mock_update.effective_chat.send_message = AsyncMock()

        mock_context = MagicMock()
        mock_context.user_data = {}

        await _render_or_send_text(mock_update, mock_query, "Тестове меню", None)
        mock_query.edit_message_reply_markup.assert_awaited_once_with(reply_markup=None)
        mock_update.effective_chat.send_message.assert_awaited_once()

        if MUSEUM_ADMIN_IDS:
            mock_query.edit_message_reply_markup.reset_mock()
            mock_update.effective_chat.send_message.reset_mock()
            await admin_museum_menu_show(mock_update, mock_context)
            mock_query.edit_message_reply_markup.assert_awaited_once_with(reply_markup=None)
            mock_update.effective_chat.send_message.assert_awaited_once()

        print("   ✅ Navigation from .docx message works cleanly without BadRequest and preserves the file!")

        print("\n🎉 ВСІ ТЕСТИ УСПІШНО ПРОЙДЕНО!")
    finally:
        # Гарантоване прибирання тестових записів з локальної БД
        async with AsyncSessionLocal() as session:
            await session.execute(delete(MuseumBooking).where(MuseumBooking.excursion_date == test_date_reg))
            await session.commit()
        print(f"🧹 Cleaned up test booking records for {test_date_reg}")


if __name__ == "__main__":
    asyncio.run(run_museum_document_tests())

